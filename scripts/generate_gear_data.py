#!/usr/bin/env python3
"""
Generate src/gear_set_data.py, the set table the app ships, from the LibSets
addon's workbook (data/gear_sets/LibSets_SetData.xlsm).

Design: the encounter log names a gear piece's set only by its set id
(PLAYER_INFO ... setId). LibSets, the addon the ESO community keeps for
exactly this mapping, publishes a workbook with every set's id, name, type
(crafted, trial, mythic, ...) and item ids; this script reads it at build
time with openpyxl and writes a Python module, so nothing is parsed at run
time and the workbook is not bundled. The CI build runs the script, so a
change here is what ships. Set names are the game's own (issue #79): the
installed client's English string table, committed as
data/client/set_names_en.json by scripts/check_client_names.py, names every
set, and the workbook's spelling differed for 26 of them on Update 51
("Perfect" for "Perfected", "Blood Spawn", three Cyrodiil sets one id out of
place). The workbook gives the set type, the one field the app reads from it
(the five-piece reading in esolog_tail.has_five_piece_bonus, the mythic one
in player_build.is_mythic_set), and the names of the template and "(OLD)"
rows the client lacks. A set the client has and the workbook does not yet is
included under its client name with no type, so it reads as a five-piece,
non-mythic set until the workbook catches up. ESO-Hub links for sets are
looked up by name (see src/ability_icons.py).
"""

import json
import os
from pathlib import Path
from typing import Dict, List, Set

from openpyxl import load_workbook

def filled(value) -> bool:
    """True when a worksheet cell holds something."""
    return value is not None and value != ''

# The game's own set names, by set id: string table 38727365 of the installed
# client's gamedata\lang\en.lang, committed as data/client/set_names_en.json by
# scripts/check_client_names.py --write-sets (CI has no client to read). They
# replace the workbook's names, which differed for 26 sets on Update 51
# ("Perfect" for "Perfected", "Blood Spawn", three Cyrodiil sets one id out of
# place). The workbook still gives each set's type, and the names of the
# template and "(OLD)" rows the client does not have.
CLIENT_NAMES_PATH = Path(__file__).resolve().parent.parent / "data" / "client" / "set_names_en.json"


def load_client_names(path: Path = CLIENT_NAMES_PATH) -> Dict[str, str]:
    """{set id: name} from the committed client snapshot; empty, with a
    warning, when it is missing."""
    if not path.is_file():
        print(f"Warning: {path} not found; set names come from the workbook only")
        return {}
    return json.loads(path.read_text(encoding="utf-8"))["sets"]


def set_names(workbook: Dict[str, str], client: Dict[str, str]) -> Dict[str, str]:
    """{set id: name} for every set either source has: the client's name
    when it has one, else the workbook's. A set only the client has (new
    since the workbook) is included under its client name."""
    names = {set_id: client.get(set_id, name) for set_id, name in workbook.items()}
    for set_id, name in client.items():
        names.setdefault(set_id, name)
    renamed = sorted((k for k in workbook if k in client and client[k] != workbook[k]), key=int)
    client_only = sorted(set(client) - set(workbook), key=int)
    print(f"Set names: {len(names)} sets, {len(renamed)} named by the client where the "
          f"workbook differs, {len(client_only)} from the client only")
    for set_id in client_only:
        print(f"  client only {set_id}: {client[set_id]} (no set type until the workbook has it)")
    return names


def extract_gear_data():
    """Extract gear set data from the XLSM file."""
    current_dir = Path(__file__).parent.parent
    excel_file_path = current_dir / "data" / "gear_sets" / "LibSets_SetData.xlsm"
    
    if not excel_file_path.exists():
        print(f"Warning: Excel file not found: {excel_file_path}")
        return None
    
    try:
        # Read the Excel file (cached values, not formulas)
        workbook = load_workbook(excel_file_path, read_only=True, data_only=True)
        # Try both old and new sheet names for compatibility
        sheet_name = 'Sets (lua data)' if 'Sets (lua data)' in workbook.sheetnames else 'Sets data'
        rows = workbook[sheet_name].iter_rows(values_only=True)
        next(rows)  # a "Last updated" row sits above the column headings
        headings = next(rows)
        # One dict per row, keyed by column heading; empty rows are dropped
        table = [dict(zip(headings, row)) for row in rows
                 if any(filled(cell) for cell in row)]
        workbook.close()
        
        print(f"Extracted {len(table)} gear sets from {excel_file_path}")
        
        # Extract the data we need
        gear_data = {
            'set_id_to_name': {},
            'set_name_to_id': {},
            'set_info': {},
            'known_item_mappings': {},
            'known_ability_mappings': {}
        }
        
        # The workbook's name, type and comment per set id
        workbook_names = {}
        workbook_info = {}
        for row in table:
            set_id = row.get('ESO ingame setId')
            set_name_en = row.get('Name EN')
            set_type = row.get('Set Type', '')
            comment = row.get('Comment', '')

            if filled(set_id) and filled(set_name_en):
                set_id_str = str(int(set_id))
                # The workbook carries Lua escapes in a few names
                # ("Mara\’s Balm", "Siegemaster'\s Focus"). The backslash is
                # never part of a set's name, and the typographic apostrophe
                # becomes the plain one every other name uses.
                workbook_names[set_id_str] = (str(set_name_en).strip().strip('"')
                                              .replace("\\", "").replace("’", "'"))
                workbook_info[set_id_str] = (str(set_type) if filled(set_type) else '',
                                             str(comment) if filled(comment) else '')

        # Names from the client, type from the workbook
        for set_id_str, set_name in set_names(workbook_names, load_client_names()).items():
            set_type, comment = workbook_info.get(set_id_str, ('', ''))
            gear_data['set_id_to_name'][set_id_str] = set_name
            gear_data['set_name_to_id'][set_name] = set_id_str
            gear_data['set_info'][set_name] = {
                'set_id': set_id_str,
                'set_type': set_type,
                'comment': comment,
                'items': [],
                'abilities': []
            }
        
        # Add known item and ability mappings (from the current fallback data)
        gear_data['known_item_mappings'] = {
            '154691': "Bahsei's Mania",
        }
        
        gear_data['known_ability_mappings'] = {
            '154691': "Bahsei's Mania",
            '66899': "Spell Power Cure", 
            '107202': "Arms of Relequen",
            '107203': "Arms of Relequen",
            '133378': "Mother Ciannait",
            '84731': "Witchmother's Potent Brew",
        }
        
        return gear_data
        
    except Exception as e:
        print(f"Error extracting gear data: {e}")
        return None

def quoted(text: str) -> str:
    """*text* as a double-quoted Python string literal. A JSON string is
    one, with quotes and backslashes escaped."""
    return json.dumps(text, ensure_ascii=False)


def generate_python_module(gear_data: Dict) -> str:
    """Generate Python module code with the extracted data."""

    code = '''"""
Auto-generated gear set data module.
Generated from LibSets_SetData.xlsm - DO NOT EDIT MANUALLY.

This module contains optimized data structures for fast gear set lookups
without requiring Excel file parsing at runtime.
"""

from typing import Dict, List, Optional, Set

# Set ID to Set Name mapping
SET_ID_TO_NAME: Dict[str, str] = {
'''
    
    # Add set ID to name mappings
    for set_id, set_name in sorted(gear_data['set_id_to_name'].items()):
        code += f'    "{set_id}": {quoted(set_name)},\n'
    
    code += '''}

# Set Name to Set ID mapping
SET_NAME_TO_ID: Dict[str, str] = {
'''
    
    # Add set name to ID mappings
    for set_name, set_id in sorted(gear_data['set_name_to_id'].items()):
        code += f'    {quoted(set_name)}: "{set_id}",\n'
    
    code += '''}

# Detailed set information
SET_INFO: Dict[str, Dict] = {
'''
    
    # Add set info
    for set_name, info in sorted(gear_data['set_info'].items()):
        code += f'    {quoted(set_name)}: {{\n'
        code += f'        "set_id": "{info["set_id"]}",\n'
        code += f'        "set_type": {quoted(info["set_type"])},\n'
        code += f'        "comment": {quoted(info["comment"])},\n'
        code += f'        "items": [],\n'
        code += f'        "abilities": []\n'
        code += f'    }},\n'
    
    code += '''}

# Known item ID to set name mappings
KNOWN_ITEM_MAPPINGS: Dict[str, str] = {
'''
    
    # Add known item mappings
    for item_id, set_name in sorted(gear_data['known_item_mappings'].items()):
        code += f'    "{item_id}": {quoted(set_name)},\n'
    
    code += '''}

# Known ability ID to set name mappings
KNOWN_ABILITY_MAPPINGS: Dict[str, str] = {
'''
    
    # Add known ability mappings
    for ability_id, set_name in sorted(gear_data['known_ability_mappings'].items()):
        code += f'    "{ability_id}": {quoted(set_name)},\n'
    
    code += '''}

def get_set_name_by_id(set_id: str) -> Optional[str]:
    """Get set name by set ID."""
    return SET_ID_TO_NAME.get(set_id)

def get_set_id_by_name(set_name: str) -> Optional[str]:
    """Get set ID by set name."""
    return SET_NAME_TO_ID.get(set_name)

def get_set_info(set_name: str) -> Optional[Dict]:
    """Get detailed set information by set name."""
    return SET_INFO.get(set_name)

def get_set_name_by_item_id(item_id: str) -> Optional[str]:
    """Get set name by item ID."""
    return KNOWN_ITEM_MAPPINGS.get(item_id)

def get_set_name_by_ability_id(ability_id: str) -> Optional[str]:
    """Get set name by ability ID."""
    return KNOWN_ABILITY_MAPPINGS.get(ability_id)

# Statistics
TOTAL_SETS = len(SET_ID_TO_NAME)
TOTAL_KNOWN_ITEMS = len(KNOWN_ITEM_MAPPINGS)
TOTAL_KNOWN_ABILITIES = len(KNOWN_ABILITY_MAPPINGS)

def get_stats() -> Dict[str, int]:
    """Get database statistics."""
    return {
        'total_sets': TOTAL_SETS,
        'total_known_items': TOTAL_KNOWN_ITEMS,
        'total_known_abilities': TOTAL_KNOWN_ABILITIES
    }
'''
    
    return code

def main():
    """Main function to generate the gear data module."""
    print("Generating gear set data module...")
    
    # Extract data from XLSM
    gear_data = extract_gear_data()
    if not gear_data:
        print("Failed to extract gear data")
        return 1
    
    # Generate Python module
    python_code = generate_python_module(gear_data)
    
    # Write to file
    output_file = Path(__file__).parent.parent / "src" / "gear_set_data.py"
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(python_code)
    
    print(f"Generated {output_file}")
    print(f"Statistics:")
    print(f"  - Total sets: {len(gear_data['set_id_to_name'])}")
    print(f"  - Known items: {len(gear_data['known_item_mappings'])}")
    print(f"  - Known abilities: {len(gear_data['known_ability_mappings'])}")
    
    return 0

if __name__ == "__main__":
    exit(main())
